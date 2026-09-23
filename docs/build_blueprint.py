from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(r"C:\Users\tingyun\Desktop\Tyflow Studio智能研发操作系统建设蓝图.docx")

def set_font(run, size=10.5, bold=False, color="000000"):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size); run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); pr.append(shd)

def cell_margin(cell, v=100, h=120):
    pr = cell._tc.get_or_add_tcPr(); mar = OxmlElement("w:tcMar"); pr.append(mar)
    for tag, val in (("top",v),("start",h),("bottom",v),("end",h)):
        n=OxmlElement(f"w:{tag}"); n.set(qn("w:w"),str(val)); n.set(qn("w:type"),"dxa"); mar.append(n)

def table_borders(table):
    e=OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        n=OxmlElement(f"w:{edge}"); n.set(qn("w:val"),"single"); n.set(qn("w:sz"),"4"); n.set(qn("w:color"),"D9D9D9"); e.append(n)
    table._tbl.tblPr.append(e)

def add_table(doc, headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; table_borders(t)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.width=Cm(widths[i]); shade(c,"17365D"); cell_margin(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run(h),9.5,True,"FFFFFF")
    for r,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            c=cells[i]; c.width=Cm(widths[i]); cell_margin(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r%2: shade(c,"F2F6FA")
            p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after=Pt(0)
            set_font(p.add_run(str(val)),9)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def bullet(doc,text,level=0):
    p=doc.add_paragraph(style="List Bullet" if level==0 else "List Bullet 2"); p.paragraph_format.space_after=Pt(3); set_font(p.add_run(text),10.5); return p

def lead(doc,label,text):
    p=doc.add_paragraph(); set_font(p.add_run(label+"  "),10.8,True,"1F4E78"); set_font(p.add_run(text),10.8); return p

doc=Document(); sec=doc.sections[0]
sec.top_margin=Cm(2.0); sec.bottom_margin=Cm(1.9); sec.left_margin=Cm(2.25); sec.right_margin=Cm(2.25)
styles=doc.styles
normal=styles["Normal"]; normal.font.name="Microsoft YaHei"; normal._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); normal.font.size=Pt(10.5); normal.paragraph_format.line_spacing=1.35; normal.paragraph_format.space_after=Pt(7)
for name,size,before,after in (("Title",24,0,18),("Heading 1",16,16,8),("Heading 2",12.5,11,5)):
    s=styles[name]; s.font.name="Microsoft YaHei"; s._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
sp=styles["Title"]._element.get_or_add_pPr(); b=sp.find(qn("w:pBdr"));
if b is not None: sp.remove(b)

p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run("Tyflow Studio 智能研发操作系统建设蓝图"),24,True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run("把个人 AI 编程能力升级为组织可治理 可复用 可度量的研发生产力"),12,False,"555555")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run("建设蓝图  2026 至 2027"),10,False,"666666")

doc.add_heading("一 蓝图结论",1)
lead(doc,"总体定位","Tyflow Studio 的目标不是再做一个代码生成工具，而是建设面向智能研发时代的前端交付操作系统。它把需求 规格 代码 测试 证据和发布决策连接为统一的可追踪对象，让 AI 在可控边界内承担更多执行工作，让人聚焦业务决策 风险判断和最终验收。")
doc.add_paragraph("未来形态将覆盖从需求进入到上线反馈的完整闭环，并逐步形成团队级质量策略 项目知识图谱 可复用测试资产和研发效能数据。每一次需求交付都不再是一次性对话，而是组织能力的增量沉淀。")
add_table(doc,["建设目标","目标状态","组织价值"],[
    ["研发提效","重复分析 编码和验证自动化","缩短需求交付周期"],
    ["质量前置","需求 计划 测试在开发前形成契约","减少后期返工和线上风险"],
    ["过程可信","每个结论绑定代码快照和证据","可审计 可复核 可问责"],
    ["知识复利","规则 组件 测试和失败经验持续沉淀","减少重复踩坑和人员依赖"],
    ["组织治理","统一门禁 分级授权和效能度量","AI 使用从个人行为升级为组织能力"],
],[3.2,6.7,5.1])

doc.add_heading("二 为什么现在建设",1)
doc.add_paragraph("AI 编程能力正在快速普及，但团队规模使用会同时放大效率和风险。没有工程化治理时，个人可以更快地产生代码，组织却更难确认需求是否理解正确 改动是否越界 测试是否可信以及最终由谁负责。")
add_table(doc,["现状变化","带来的机会","同步出现的风险"],[
    ["代码生成速度提升","压缩常规实现时间","错误也可以更快扩散"],
    ["AI 可跨文件执行","承担复杂改造和重复工作","范围漂移更难人工察觉"],
    ["自动化测试可由 AI 生成","提升覆盖和验证效率","自证式测试与假证据"],
    ["多 Agent 可以并行","提高复杂需求吞吐","冲突 上下文分叉和责任不清"],
    ["知识可被模型调用","经验复用成本降低","规则分散 版本冲突和过期"],
],[3.5,5.7,5.8])
lead(doc,"战略判断","下一阶段的竞争力不在于谁能调用更多模型，而在于谁能把模型能力装进稳定流程，建立统一事实源 机器门禁 证据链和持续学习机制。")

doc.add_heading("三 目标形态",1)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("业务意图  →  可执行规格  →  受控开发  →  独立验证  →  人工决策  →  发布反馈  →  组织学习"),11,True,"1F4E78")
doc.add_paragraph("目标系统由四层组成。各层职责独立，通过版本化契约连接，既能整体运行，也允许能力被其他研发工具复用。")
add_table(doc,["层级","核心对象","主要能力"],[
    ["体验与协同层","需求 计划 门禁 验收 审计","统一工作台 人工确认 风险展示 决策记录"],
    ["流程编排层","任务 状态 快照 结果 契约","流程路由 结果复用 失败恢复 并行调度"],
    ["专业能力层","Tyflow 测试 Skill 代码审查 安全扫描","需求规格 开发执行 独立测试 专项评审"],
    ["知识与数据层","项目规则 组件资产 错题 证据 指标","事实源 版本管理 经验沉淀 效能分析"],
],[3.0,5.2,6.8])

doc.add_heading("四 核心能力蓝图",1)
add_table(doc,["能力域","V1 已具备","目标能力"],[
    ["需求工程","需求录入 计划确认 显式追踪","多来源需求解析 冲突检测 规格版本治理"],
    ["开发执行","冻结版本 受控执行 差异审查","Worktree 隔离 多 Agent 编排 依赖影响分析"],
    ["质量工程","三道硬闸 P0 P1 证据 独立测试","风险自适应测试 部分结果复用 缺陷预测"],
    ["知识工程","Tyflow 规则 项目错题 组件复用","项目知识图谱 规则生命周期 自动回收过期经验"],
    ["交付治理","人工验收 提交推送 结果留档","多级审批 发布策略 变更风险评分"],
    ["效能度量","任务日志 测试结果 快照信息","交付周期 返工 Token 质量收益和趋势分析"],
],[3.0,5.8,6.2])

doc.add_heading("五 质量与治理底座",1)
doc.add_heading("三道硬闸",2)
for x in [
    "追踪闸：核心需求必须连接计划步骤 预计文件 测试用例和验证证据。",
    "证据闸：P0 P1 通过必须有机器或受控人工证据，文字结论不能放行。",
    "差异闸：实际改动必须落在计划范围内，测试阶段不得修改业务代码后直接放行。",
]: bullet(doc,x)
doc.add_heading("分级信任",2)
doc.add_paragraph("按任务风险和操作类型分配自主度。低风险重复工作可自动执行，高风险权限 数据写入 公共组件和跨项目协议必须增加人工确认与独立审查。信任等级由历史结果提升，也可因失败自动下降。")
doc.add_heading("证据优先",2)
doc.add_paragraph("所有放行结论绑定需求版本 计划哈希 代码快照 测试 Skill 版本 环境和原始证据。结果可以被复用，但只能在身份和范围完全一致时复用。")

doc.add_heading("六 典型使用场景",1)
add_table(doc,["场景","未来工作方式","关键收益"],[
    ["常规页面需求","自动生成规格 计划和测试 人工确认后流水线执行","缩短交付周期 降低理解返工"],
    ["线上 Bug","复现证据驱动最小修复 自动回归并保留失败历史","提高响应速度 防止修复扩散"],
    ["接口联调","UI 契约先行 单一映射点 真实数据读回验证","减少字段语义错误和页面返工"],
    ["跨项目复用","自动检索组件和相似实现 评估版本与适配风险","提高复用率 避免重复建设"],
    ["高风险改造","Worktree 隔离 独立审查 权限与写入专项测试","控制影响面 提高上线信心"],
    ["新人接入","项目规则 错题和组件模式按任务自动加载","降低学习成本 缩短熟悉周期"],
],[3.0,7.2,4.8])

doc.add_heading("七 分阶段建设路线",1)
add_table(doc,["阶段","时间建议","建设重点","阶段成果"],[
    ["第一阶段 可信闭环","当前至 2 个月","真实项目试点 误报治理 Worktree 隔离 部分测试复用","3 类需求稳定闭环 形成基线指标"],
    ["第二阶段 团队协同","3 至 6 个月","项目接入体检 审计时间线 多级审批 CI 集成 团队规则中心","覆盖多个前端项目 形成统一交付规范"],
    ["第三阶段 智能运营","6 至 12 个月","知识图谱 风险预测 自适应测试 效能看板 经验自动沉淀","从流程工具升级为研发运营平台"],
],[2.8,2.8,6.2,4.2])
doc.add_heading("阶段一重点验收",2)
for x in [
    "小型 Bug 标准 L2 页面和接口权限类需求均能稳定闭环。",
    "计划外改动 证据缺失和测试修改业务代码能够被稳定拦截。",
    "建立交付周期 人工确认 返工 测试重复 Token 和缺陷逃逸的基线。",
    "试点团队能够独立使用，并形成 V1.1 优先级清单。",
]: bullet(doc,x)

doc.add_heading("八 度量体系",1)
add_table(doc,["维度","核心指标","判断问题"],[
    ["效率","需求周期 AI 执行耗时 人工确认次数 重复命令减少量","是否真正更快"],
    ["质量","返工次数 门禁拦截数 缺陷逃逸率 回归失败率","是否真正更稳"],
    ["成本","Token 消耗 测试耗时 结果复用率 人工投入","效率收益是否覆盖成本"],
    ["复用","组件复用率 测试资产复用率 规则命中率","是否形成组织复利"],
    ["信任","证据完整率 人工抽查通过率 误报和漏报","团队是否敢用"],
    ["治理","高风险审批覆盖率 审计完整率 规则过期率","规模化后是否可控"],
],[2.8,7.0,5.2])

doc.add_page_break(); doc.add_heading("九 组织落地模式",1)
doc.add_paragraph("建议采用小核心团队加项目共建人的方式推进。平台团队维护流程 契约和共性能力，项目代表维护项目规则 测试数据和验收基线，研发与测试共同定义放行标准。")
add_table(doc,["角色","主要责任","阶段投入"],[
    ["产品或业务负责人","确认需求目标 业务语义和最终验收","关键节点参与"],
    ["平台建设负责人","维护 Studio 编排 门禁 契约和版本","持续投入"],
    ["项目研发代表","维护项目规范 组件模式和工程命令","试点期重点投入"],
    ["测试代表","定义关键用例 独立 Oracle 和抽查策略","试点期重点投入"],
    ["安全与发布角色","定义高风险操作和发布审批边界","按需参与"],
],[3.1,7.1,4.8])

doc.add_heading("十 风险与控制原则",1)
add_table(doc,["风险","控制措施"],[
    ["过度自动化导致责任模糊","保留业务语义 高风险操作和最终发布的人类决策权"],
    ["规则过多拖慢研发","按风险分级执行 记录误报并持续删除低价值门禁"],
    ["AI 自证导致假通过","独立测试提供者 结构化证据 代码快照和人工抽查"],
    ["多项目规则冲突","目标仓库真实约束优先 规则版本化并公开冲突"],
    ["敏感信息进入模型或报告","最小化上下文 脱敏 受控证据目录和安全扫描"],
    ["成本不可控","记录 Token 与运行时间 推动结果复用和风险自适应测试"],
],[4.0,11.0])

doc.add_heading("十一 预期终局",1)
lead(doc,"终局形态","每一个研发任务都有明确业务意图 可执行规格 受控变更 独立证据和可追溯决策；每一次失败都会更新项目知识，每一次成功都能沉淀为下一次默认能力。")
doc.add_paragraph("届时，Tyflow Studio 不只是提升单个开发者效率，而是成为连接产品 研发 测试和交付治理的智能研发基础设施。团队可以在不牺牲质量和责任边界的前提下逐步提升 AI 自主度，把研发能力从个人经验驱动升级为组织系统驱动。")

doc.add_heading("十二 建议决策",1)
doc.add_paragraph("建议将当前版本确认为 V1 试运行基线，选择 1 至 2 个前端项目开展四周试点；试点期维持现有发布审批制度，以效率 质量 成本和信任数据评审是否进入团队推广阶段。")
for x in [
    "批准三类代表性需求进入真实试点。",
    "指定平台 研发和测试共建人。",
    "允许采集脱敏后的交付效率 质量和 Token 指标。",
    "四周后进行阶段评审，确定 V1.1 和团队化建设投入。",
]: bullet(doc,x)

doc.add_page_break(); doc.add_heading("附录 当前起点",1)
doc.add_paragraph("当前 V1 已具备需求到提交的主流程 三道硬闸 测试结果复用 Tyflow 独立验证和结构化测试契约，并完成完整自动化回归。")
add_table(doc,["能力","状态"],[
    ["需求 计划 测试显式追踪","已具备"],["计划文件范围与差异审查","已具备"],["P0 P1 结构化证据","已具备"],["测试前后代码防篡改","已具备"],["TestRequest 与 TestResult V1","已具备"],["完整自动化回归","10 个套件通过 102 项通过 1 项跳过 0 项失败"],
],[7.5,7.5])

for s in doc.sections:
    p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run("Tyflow Studio 智能研发操作系统建设蓝图"),8.5,False,"777777")
doc.core_properties.title="Tyflow Studio 智能研发操作系统建设蓝图"; doc.core_properties.subject="智能研发平台建设路线"; doc.core_properties.author=""
doc.save(OUT); print(OUT)
